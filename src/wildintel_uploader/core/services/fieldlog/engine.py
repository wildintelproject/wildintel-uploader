"""The field-records engine — the FileTimestamps tool (Tatra and Doñana camera-trap projects) brought in as it is,
with what the app needs around it: errors raised instead of exiting, the visits carrying who, which card and the notes
(for the FieldworkLog), the deployments the two visits they span, and the Epicollect columns overridable from the settings.

It turns the Epicollect CSV download of the last field trip and the fieldwork datasheet (.xlsx) into the updated
datasheet and the deployments' start and end — every time in UTC, the time of the Epicollect entry (created_at).
See fieldlog_service for what the app does with it."""
import collections
import csv
import datetime as dt
import io
import os
import re
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter


class FieldlogError(Exception):
    """The inputs can't be read as an Epicollect download and a fieldwork datasheet."""

# ==============================================================================
# Configuration: the only place where names are defined
# ==============================================================================

# Columns of the fieldwork datasheet (first worksheet, row 1).
SHEET_HEADERS = ['Trappers', 'Session', 'Date', 'Time', 'DateTime', 'Site', 'Camera', 'Action',
                 'Status', 'CardBatteriesOK', 'NumPhotos', 'SD', 'DateTimeOK', 'Notes',
                 'Longitude', 'Latitude']

# datasheet column  <-  Epicollect column (matched by name, tolerant to the numeric prefix, case and spaces)
EPI_TO_SHEET = collections.OrderedDict([
    ('Trappers',        '1_Responsable'),
    ('Session',         '2_Collection'),
    ('Site',            '4_Location'),
    ('Camera',          '5_Camera'),
    ('Action',          '6_Action'),
    ('Status',          '7_Status_of_the_came'),
    ('CardBatteriesOK', '8_Card_Batteries'),
    ('NumPhotos',       '9_Number_of_pictures'),
    ('SD',              '10_SD'),
    ('DateTimeOK',      '11_Date_and_hour'),
    ('Notes',           '12_Observation'),
    ('Longitude',       'long_13_Coordenates'),
    ('Latitude',        'lat_13_Coordenates'),
])
# Camtrap DP fields of a deployment the field app may also ask for (setupBy, cameraID and deploymentComments come from
# Trappers, Camera and Notes). Each is read from the column of its own name, or the one the settings map it to, when the
# field survey has it; kept in the field history under the same name, so that a later trip still has it.
EXTRA_FIELDS = ['cameraModel', 'cameraDelay', 'cameraHeight', 'cameraDepth', 'cameraTilt', 'cameraHeading',
                'detectionDistance', 'timestampIssues', 'baitUse', 'featureType', 'habitat', 'deploymentGroups',
                'deploymentTags']
EPI_CREATED = 'created_at'      # UTC timestamp of the entry: gives Date and Time of the visit (as in R0012 and R0013)
EPI_DATE = '3_Date'             # typed date, only used as a cross check
EPI_HOUR = '14_Hour'            # local hour, only used with -timesource hour

# Header of the final log
LOG_HEADER = ['Deployment', 'StartDate', 'StartTime', 'EndDate', 'EndTime', ' ',
              'Deployment_Or', 'StartDate_Or', 'StartTime_Or', 'EndDate_Or', 'EndTime_Or']

END_ACTIONS = {'desinstallation', 'uninstallation', 'removal', 'd'}   # camera taken away: no interval after it
LOST_ACTIONS = {'stolen'}                                              # no interval ending in it
NOT_VISITED_STATUS = {'not revised', 'no revisada', 'no revisado'}     # nothing was collected at that visit
DEFAULT_SITE_PREFIX = 'TATR_'

# ------------------------------------------------------------------------------
# Donana files (Spanish headers, R/I/D/ROBADA). The app recognises them, works on a translated copy in
# memory and can write the updated datasheet in its original format, in the standard format, or both.
# ------------------------------------------------------------------------------
ES_SHEET_HEADERS = collections.OrderedDict([
    ('Tramperos', 'Trappers'), ('Sesion', 'Session'), ('Fecha', 'Date'), ('Hora', 'Time'),
    ('FechaHora', 'DateTime'), ('Sitio', 'Site'), ('Camara', 'Camera'), ('Accion', 'Action'),
    ('Estado', 'Status'), ('CardPilasOK', 'CardBatteriesOK'), ('NumFotos', 'NumPhotos'), ('SD', 'SD'),
    ('FechaHoraOK', 'DateTimeOK'), ('Observaciones', 'Notes'), ('Longitud', 'Longitude'), ('Latitud', 'Latitude'),
])
ES_EPI_HEADERS = {
    '1_Tramperos': '1_Responsable', '2_Sesion': '2_Collection', '3_Fecha': '3_Date', '4_Sitio': '4_Location',
    '5_Camara': '5_Camera', '6_Accion': '6_Action', '7_Estado': '7_Status_of_the_came',
    '8_CardPilasOK': '8_Card_Batteries', '9_NumFotos': '9_Number_of_pictures',
    '11_FechaHoraOK': '11_Date_and_hour', '12_Observaciones': '12_Observation',
    'lat_13_Localizacin_actua': 'lat_13_Coordenates', 'long_13_Localizacin_actua': 'long_13_Coordenates',
    '14_Hora': '14_Hour',
}
ES_ACTIONS = {'R': 'Revision', 'I': 'Installation', 'D': 'Desinstallation', 'ROBADA': 'Stolen'}
ES_ACTIONS_BACK = {v: k for k, v in ES_ACTIONS.items()}
ES_SITE_PREFIX = 'DONA_'


# ==============================================================================
# Small helpers
# ==============================================================================

class Report(object):
    def __init__(self):
        self.lines = []
        self.warnings = 0

    def info(self, msg):
        self.lines.append(msg)

    def warn(self, msg):
        self.warnings += 1
        self.lines.append('WARNING: ' + msg)

    def text(self):
        return '\n'.join(self.lines) + '\n'


def norm(name):
    """Normalise a column name: lower case, no spaces, no leading '12_' style prefix."""
    n = re.sub(r'\s+', '', str(name)).lower()
    return re.sub(r'^\d+_', '', n)


def find_column(headers, wanted):
    for h in headers:
        if h == wanted:
            return h
    nw = norm(wanted)
    for h in headers:
        if norm(h) == nw:
            return h
    return None


def parse_date(s):
    s = str(s).strip()
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError('unreadable date: %r' % s)


def parse_created(s):
    """Epicollect timestamps look like 2026-09-04T10:16:31.000Z (UTC). Returns a naive UTC datetime."""
    s = str(s).strip().replace('Z', '+00:00')
    d = dt.datetime.fromisoformat(s)
    if d.tzinfo is not None:
        d = (d - d.utcoffset()).replace(tzinfo=None)
    return d.replace(microsecond=0)


def parse_time_str(s):
    m = re.match(r'^\s*(\d{1,2}):(\d{2})(?::(\d{2}))?', str(s))
    if not m:
        raise ValueError('unreadable time: %r' % s)
    return dt.time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))


def cell_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=v)).date()
    return parse_date(v)


def cell_time(v):
    if isinstance(v, dt.datetime):
        return v.time().replace(microsecond=0)
    if isinstance(v, dt.time):
        return v.replace(microsecond=0)
    if isinstance(v, (int, float)):
        secs = int(round((v % 1) * 86400))
        return (dt.datetime.min + dt.timedelta(seconds=secs)).time()
    return parse_time_str(v)


def session_number(s):
    m = re.search(r'(\d+)', str(s))
    return int(m.group(1)) if m else None


def session_name(n):
    return 'R%04d' % n


def to_int(v):
    if v is None or str(v).strip() == '':
        return None
    try:
        return int(float(str(v).replace(',', '.')))
    except ValueError:
        return None


def to_float(v):
    if v is None or str(v).strip() == '':
        return None
    try:
        return float(str(v).replace(',', '.'))
    except ValueError:
        return None


# ==============================================================================
# Reading the inputs
# ==============================================================================

def read_epicollect(path, report, overrides=None):
    """overrides: {datasheet column: Epicollect column} that replace the default names of EPI_TO_SHEET (and of
    created_at, 3_Date and 14_Hour, by those names) for forms whose questions are named otherwise."""
    overrides = {k: v for k, v in (overrides or {}).items() if v and str(v).strip()}
    with io.open(path, 'r', encoding='utf-8-sig', newline='') as fh:
        text = fh.read()
    first = text.split('\n', 1)[0]
    delim = max([',', ';', '\t'], key=lambda d: first.count(d))
    reader = csv.DictReader(io.StringIO(text, newline=''), delimiter=delim)
    headers = reader.fieldnames or []
    reader = list(reader)
    if find_column(headers, '2_Collection') is None and find_column(headers, '2_Sesion') is not None:
        ren = dict((find_column(headers, k), v) for k, v in ES_EPI_HEADERS.items() if find_column(headers, k))
        headers = [ren.get(h, h) for h in headers]
        reader = [dict((ren.get(k, k), v) for k, v in r.items()) for r in reader]
        act = find_column(headers, '6_Action')
        for r in reader:
            if act and r.get(act):
                r[act] = _es_action(r[act])
        report.info('Epicollect format recognised: Donana (Spanish headers). Translated in memory.')
    cols = {}
    missing = []
    for sheet_col, epi_col in EPI_TO_SHEET.items():
        epi_col = overrides.get(sheet_col, epi_col)
        found = find_column(headers, epi_col)
        if found is None:
            missing.append(epi_col)
        cols[sheet_col] = found
    for extra in (EPI_CREATED, EPI_DATE):
        found = find_column(headers, overrides.get(extra, extra))
        if found is None:
            missing.append(extra)
        cols[extra] = found
    cols[EPI_HOUR] = find_column(headers, overrides.get(EPI_HOUR, EPI_HOUR))
    for field in EXTRA_FIELDS:
        cols[field] = find_column(headers, overrides.get(field, field))
    found_extra = [f for f in EXTRA_FIELDS if cols[f]]
    if found_extra:
        report.info('Camtrap DP fields in the field survey: %s' % ', '.join(found_extra))
    if missing:
        raise FieldlogError('The Epicollect file does not have these columns: %s. '
                         'Columns found: %s. '
                         'Check the question names of the form, or map them in Settings > Field records.'
                         % (', '.join(missing), ', '.join(headers)))
    rows = []
    for r in reader:
        if not any((v or '').strip() for v in r.values()):
            continue
        rows.append(r)
    report.info('Epicollect file: %s (%d entries)' % (os.path.basename(path), len(rows)))
    return rows, cols


def _es_session(v):
    return 'R%04d' % int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v


def _es_site(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return '%s%04d' % (ES_SITE_PREFIX, int(v))
    if isinstance(v, str):
        m = re.match(r'^\s*(\d+)\s*([bB]?)\s*$', v)
        if m:
            return '%s%04d%s' % (ES_SITE_PREFIX, int(m.group(1)), '_B' if m.group(2) else '')
    return v


def _es_action(v):
    return ES_ACTIONS.get(v.strip().upper(), v) if isinstance(v, str) else v


def sheet_headers(ws):
    return [str(c.value).strip() if c.value is not None else '' for c in ws[1]]


def standardize_sheet(wb, report):
    """If the first datasheet is in the Donana (Spanish) format, translate it IN MEMORY to the standard one.
    Returns 'Donana' or None (already standard)."""
    for ws in wb.worksheets:
        heads = sheet_headers(ws)
        if 'Session' in heads and 'Site' in heads and 'Camera' in heads:
            return None
        if 'Sesion' in heads and 'Sitio' in heads:
            missing = [h for h in ES_SHEET_HEADERS if h not in heads]
            if missing:
                raise FieldlogError('The datasheet looks like the Donana format but lacks these columns: %s'
                                 % ', '.join(missing))
            col = {h: i + 1 for i, h in enumerate(heads) if h}
            for old, new in ES_SHEET_HEADERS.items():
                ws.cell(1, col[old]).value = new
            for r in range(2, ws.max_row + 1):
                if all(ws.cell(r, c).value in (None, '') for c in range(1, len(heads) + 1)):
                    continue
                for name, fn in (('Sesion', _es_session), ('Sitio', _es_site), ('Accion', _es_action)):
                    c = ws.cell(r, col[name])
                    new = fn(c.value)
                    if new != c.value:
                        c.value = new
            report.info('Datasheet format recognised: Donana (Spanish headers, actions R/I/D/ROBADA). '
                        'It is translated in memory; your original file is not changed.')
            return 'Donana'
    return None


def find_original_sheet(wb):
    """The Donana datasheet as it is in the file, with column positions under the standard names."""
    for ws in wb.worksheets:
        heads = sheet_headers(ws)
        if 'Sesion' in heads and 'Sitio' in heads:
            return ws, {ES_SHEET_HEADERS[h]: i + 1 for i, h in enumerate(heads) if h in ES_SHEET_HEADERS}
    raise FieldlogError('Original datasheet not found.')


def rows_to_original(new, ws_o, colidx_o):
    """New rows written the way the original datasheet writes them (same style as its last row)."""
    last = last_data_row(ws_o, colidx_o)
    sess_num = isinstance(ws_o.cell(last, colidx_o['Session']).value, (int, float))
    site_num = isinstance(ws_o.cell(last, colidx_o['Site']).value, (int, float))
    out = []
    for row in new:
        d = dict(row)
        if sess_num:
            d['Session'] = row['snum']
        m = re.match(r'^%s0*(\d+)$' % ES_SITE_PREFIX, row['Site'])
        if site_num and m:
            d['Site'] = int(m.group(1))
        d['Action'] = ES_ACTIONS_BACK.get(row['Action'], row['Action'])
        out.append(d)
    return out


def find_sheets(wb):
    data = None
    for ws in wb.worksheets:
        heads = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
        if data is None and 'Session' in heads and 'Site' in heads and 'Camera' in heads:
            data = (ws, heads)
    if data is None:
        raise FieldlogError('No worksheet with the field data columns (Session, Site, Camera...) in row 1.')
    ws, heads = data
    missing = [h for h in SHEET_HEADERS if h not in heads]
    if missing:
        raise FieldlogError('The field data sheet lacks these columns: %s' % ', '.join(missing))
    colidx = {h: heads.index(h) + 1 for h in heads if h}
    return ws, colidx


def read_sheet_visits(ws, colidx, report):
    visits = []
    for r in range(2, ws.max_row + 1):
        g = lambda name: ws.cell(r, colidx[name]).value
        site, sess = g('Site'), g('Session')
        if site in (None, '') and sess in (None, ''):
            continue
        try:
            when = dt.datetime.combine(cell_date(g('Date')), cell_time(g('Time')))
        except Exception as e:
            report.warn('datasheet row %d (%s %s): date/time not readable (%s), row ignored' % (r, sess, site, e))
            continue
        np_ = g('NumPhotos')
        try:
            np_ = float(np_) if np_ not in (None, '') else None
        except (TypeError, ValueError):
            np_ = None
        visits.append(dict(row=r, site=str(site).strip(), session=str(sess).strip(),
                           snum=session_number(sess), cam=to_int(g('Camera')), when=when,
                           action=str(g('Action') or '').strip(), nphotos=np_,
                           status=str(g('Status') or '').strip(), trappers=_text(g('Trappers')),
                           sd=_text(g('SD')), notes=_text(g('Notes')),
                           lat=to_float(g('Latitude')), lon=to_float(g('Longitude')),
                           extra={f: _text(g(f)) for f in EXTRA_FIELDS if f in colidx}))
    return visits


def site_prefix(visits):
    c = collections.Counter(re.match(r'^(.*?)\d+$', v['site']).group(1)
                            for v in visits if re.match(r'^(.*?)\d+$', v['site']))
    return c.most_common(1)[0][0] if c else DEFAULT_SITE_PREFIX


# ==============================================================================
# Epicollect > new datasheet rows
# ==============================================================================

def epicollect_to_rows(epi_rows, cols, prefix, time_source, report):
    out = []
    legacy = 0
    for i, r in enumerate(epi_rows, start=2):
        get = lambda k: (r.get(cols[k]) or '').strip()
        label = 'Epicollect line %d' % i
        snum = session_number(get('Session'))
        loc = get('Site')
        if snum is None or not loc:
            legacy += 1
            continue
        site = loc if not re.match(r'^\d+$', loc) else '%s%04d' % (prefix, int(loc))
        try:
            created = parse_created(get(EPI_CREATED))
        except Exception as e:
            report.warn('%s (%s %s): %s, entry ignored' % (label, session_name(snum), site, e))
            continue
        when = created
        typed = None
        if get(EPI_DATE):
            try:
                typed = parse_date(get(EPI_DATE))
            except ValueError:
                pass
        if time_source == 'hour' and typed is not None and cols[EPI_HOUR] and get(EPI_HOUR):
            when = dt.datetime.combine(typed, parse_time_str(get(EPI_HOUR)))
        out.append(dict(typed=typed, created=created,
            Trappers=get('Trappers') or None, Session=session_name(snum), snum=snum,
            when=when, Site=site, Camera=to_int(get('Camera')), Action=get('Action') or None,
            Status=get('Status') or None, CardBatteriesOK=get('CardBatteriesOK') or None,
            NumPhotos=to_int(get('NumPhotos')), SD=get('SD') or None, DateTimeOK=get('DateTimeOK') or None,
            Notes=get('Notes') or None, Longitude=to_float(get('Longitude')), Latitude=to_float(get('Latitude')),
            extra={f: (r.get(cols[f]) or '').strip() or None for f in EXTRA_FIELDS if cols.get(f)}))
    if legacy:
        report.info('Epicollect entries without collection number or location (old form versions), ignored: %d' % legacy)
    return out


def select_new(rows, visits, report):
    last = max(v['snum'] for v in visits if v['snum'] is not None)
    have = set((v['snum'], v['site'], v['cam'], v['action'].lower()) for v in visits)
    new, seen_old = [], 0
    for r in rows:
        if r['snum'] < last:
            continue
        key = (r['snum'], r['Site'], r['Camera'], (r['Action'] or '').lower())
        if key in have:
            seen_old += 1
        else:
            new.append(r)
    new.sort(key=lambda r: (r['snum'], r['Site'], r['when']))
    for r in new:
        if r['typed'] is not None and r['typed'] != r['created'].date():
            report.warn('%s %s: typed date %s differs from the entry timestamp date %s (UTC); the timestamp was used'
                        % (r['Session'], r['Site'], r['typed'], r['created'].date()))
    report.info('Last session already in the datasheet: %s' % session_name(last))
    report.info('Entries of %s or later already in the datasheet: %d' % (session_name(last), seen_old))
    if new:
        sess = sorted(set(r['snum'] for r in new))
        report.info('New rows to add: %d (sessions %s)' % (len(new), ', '.join(session_name(s) for s in sess)))
        if sess[0] > last + 1:
            report.warn('Session numbers jump from %s to %s' % (session_name(last), session_name(sess[0])))
    else:
        report.info('New rows to add: 0 (the datasheet is already up to date)')
    return new


def new_to_visits(new):
    return [dict(row=None, site=r['Site'], session=r['Session'], snum=r['snum'], cam=r['Camera'],
                 when=r['when'], action=(r['Action'] or ''), nphotos=r['NumPhotos'],
                 status=(r['Status'] or ''), trappers=r['Trappers'], sd=r['SD'], notes=r['Notes'],
                 lat=r['Latitude'], lon=r['Longitude'], extra=dict(r.get('extra') or {}), new=True) for r in new]


def _text(v):
    """A cell as text, None when blank."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = str(v).strip()
    return v or None


# ==============================================================================
# Checks on the visit history
# ==============================================================================

def check_history(visits, check_sessions, report):
    by_site = collections.defaultdict(list)
    for v in visits:
        by_site[v['site']].append(v)
    for s in check_sessions:
        cur = [v for v in visits if v['snum'] == s]
        prev = set(v['site'] for v in visits if v['snum'] == s - 1)
        miss = sorted(prev - set(v['site'] for v in cur))
        if miss:
            # Not a problem: a camera left on the trip before and not visited on this one is still out, its deployment
            # open until a later trip collects its card (the page lists the open deployments).
            report.info('%s: sites visited in %s and not on this trip, their cameras still out: %s'
                        % (session_name(s), session_name(s - 1), ', '.join(miss)))
        cnt = collections.Counter((v['site'], v['action'].lower()) for v in cur)
        for (site, act), n in sorted(cnt.items()):
            if n > 1 and act == 'revision':
                report.warn('%s: %s has %d "Revision" entries in the same session' % (session_name(s), site, n))
        for (site, act), n in sorted(cnt.items()):
            if n > 1 and act == 'installation':
                report.warn('%s: %s has %d "Installation" entries in the same session (typo in the site number?)'
                            % (session_name(s), site, n))
        for v in cur:
            older = [e for e in by_site[v['site']] if e['snum'] < v['snum'] and e['when'] >= v['when']]
            if older:
                report.warn('%s %s: visit time %s is not after the visit of %s (%s)'
                            % (v['session'], v['site'], v['when'], older[0]['session'], older[0]['when']))
            if v['action'].lower() == 'revision' and v['nphotos'] is not None and v['nphotos'] <= 0:
                report.warn('%s %s: Revision with 0 photos; no deployment will be produced for it'
                            % (v['session'], v['site']))
            if v['cam'] is None:
                report.warn('%s %s: camera number is empty' % (v['session'], v['site']))
        nocount = sorted(v['site'] for v in cur if closes_collection(v) and v['nphotos'] is None)
        if nocount:
            report.info('%s: %d visit(s) without photo count (card full or count forgotten), counted as the end '
                        'of the collection: %s' % (session_name(s), len(nocount), ', '.join(nocount)))
        skipped = sorted(v['site'] for v in cur if v['status'].lower() in NOT_VISITED_STATUS)
        if skipped:
            report.info('%s: %d visit(s) with status "%s", ignored (no card collected): %s'
                        % (session_name(s), len(skipped), 'Not revised', ', '.join(skipped)))


# ==============================================================================
# Deployment intervals > FileTimestampLog
# ==============================================================================

def closes_collection(v):
    """True if this visit is the end of a collection, so a deployment ends here.
    Any Revision or Desinstallation counts, also when the number of photos is blank (card full, or the
    fieldworker forgot it, or text such as 'CF'). Not counted: an explicit 0 photos (nothing to timestamp),
    an Installation (no card taken), a stolen camera, and visits where nothing was collected (Not revised)."""
    act = v['action'].lower()
    if act in LOST_ACTIONS or not (act == 'revision' or act in END_ACTIONS):
        return False
    if v.get('status', '').lower() in NOT_VISITED_STATUS:
        return False
    return v['nphotos'] is None or v['nphotos'] > 0


def build_intervals(visits, keep_seconds, report, check_sessions):
    """Each site's visits are sorted by time and every pair of consecutive visits is a deployment, named after
    the session of the later visit minus one. The later visit must close a collection (see closes_collection)."""
    def key(v):
        return v['when'] if keep_seconds else v['when'].replace(second=0)

    by_site = collections.defaultdict(list)
    for v in visits:
        if v.get('status', '').lower() in NOT_VISITED_STATUS:
            continue
        by_site[v['site']].append(v)
    out = []
    for site in sorted(by_site):
        g = sorted(by_site[site], key=key, reverse=True)
        for i in range(len(g) - 1):
            later, earlier = g[i], g[i + 1]
            if earlier['action'].lower() in END_ACTIONS:
                continue
            if not closes_collection(later):
                continue
            if (later['snum'] in check_sessions and earlier['snum'] != later['snum'] - 1
                    and earlier['action'].lower() != 'installation'):
                report.warn('%s %s: previous visit was in %s, not in the session before; the deployment '
                            'spans several sessions' % (later['session'], site, earlier['session']))
            out.append(dict(dep=later['snum'] - 1, site=site, start=key(earlier), end=key(later),
                            start_visit=earlier, end_visit=later))
    return out


def fmt_d(d):
    return d.strftime('%Y:%m:%d')


def fmt_t(d, keep_seconds):
    return '%d:%02d:%02d' % (d.hour, d.minute, d.second if keep_seconds else 0)


def write_log(path, intervals, keep_seconds):
    with io.open(path, 'w', encoding='ascii', newline='') as fh:
        fh.write(','.join(LOG_HEADER) + '\r\n')
        for iv in intervals:
            name = '%s-%s' % (session_name(iv['dep']), iv['site'])
            vals = [name, fmt_d(iv['start']), fmt_t(iv['start'], keep_seconds),
                    fmt_d(iv['end']), fmt_t(iv['end'], keep_seconds), '',
                    name, fmt_d(iv['start']), fmt_t(iv['start'], keep_seconds),
                    fmt_d(iv['end']), fmt_t(iv['end'], keep_seconds)]
            fh.write(','.join(vals) + '\r\n')


# ==============================================================================
# Updating the workbook (keeps layout; only appends rows)
# ==============================================================================

def last_data_row(ws, colidx):
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, colidx['Site']).value not in (None, ''):
            last = r
    return last


def extra_columns(ws, colidx, new):
    """The Camtrap DP fields the new rows bring that the datasheet has no column for yet: added at the end of row 1,
    named as Camtrap DP names them, so that the field history keeps them for the trips to come."""
    present = {f for row in new for f, v in (row.get('extra') or {}).items() if v not in (None, '')}
    heads = sheet_headers(ws)
    for field in EXTRA_FIELDS:
        if field in colidx or field not in present:
            continue
        if field in heads:
            colidx[field] = heads.index(field) + 1
            continue
        col = max(len(heads), max(colidx.values())) + 1
        ws.cell(1, col).value = field
        ws.cell(1, col)._style = copy(ws.cell(1, col - 1)._style)
        heads = sheet_headers(ws)
        colidx[field] = col


def append_rows(ws, colidx, new):
    extra_columns(ws, colidx, new)
    last = last_data_row(ws, colidx)
    ncols = max(colidx.values())
    for k, row in enumerate(new, start=1):
        r = last + k
        for c in range(1, ncols + 1):
            ws.cell(r, c)._style = copy(ws.cell(last, c)._style)

        def put(name, val):
            ws.cell(r, colidx[name]).value = val

        for name in ('Trappers', 'Session', 'Site', 'Camera', 'Action', 'Status', 'CardBatteriesOK',
                     'NumPhotos', 'SD', 'DateTimeOK', 'Notes', 'Longitude', 'Latitude'):
            put(name, row[name])
        for field, value in (row.get('extra') or {}).items():
            if field in colidx and value not in (None, ''):
                put(field, value)
        put('Date', dt.datetime.combine(row['when'].date(), dt.time(0, 0)))
        put('Time', row['when'].time())
        ws.cell(r, colidx['Date']).number_format = 'dd/mm/yyyy;@'
        ws.cell(r, colidx['Time']).number_format = 'hh:mm:ss;@'
        dcol = get_column_letter(colidx['Date'])
        tcol = get_column_letter(colidx['Time'])
        put('DateTime', '=%s%d+%s%d' % (dcol, r, tcol, r))
        ws.cell(r, colidx['DateTime']).number_format = 'dd/mm/yyyy hh:mm:ss'
    newlast = last + len(new)
    if ws.auto_filter and ws.auto_filter.ref:
        m = re.match(r'^([A-Z]+)(\d+):([A-Z]+)(\d+)$', ws.auto_filter.ref)
        if m and int(m.group(4)) < newlast:
            ws.auto_filter.ref = '%s%s:%s%d' % (m.group(1), m.group(2), m.group(3), newlast)
    return newlast


# ==============================================================================
# Main workflow
# ==============================================================================

def unique_path(path):
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    stamp = dt.datetime.now().strftime('%Y%m%d_%H%M%S')
    return '%s_%s%s' % (base, stamp, ext)


def new_sheet_name(sheet_path, first, last, suffix=''):
    stem, ext = os.path.splitext(os.path.basename(sheet_path))
    rng = '%s-%s' % (session_name(first), session_name(last))
    if re.search(r'R\d{4}-R\d{4}', stem):
        stem = re.sub(r'R\d{4}-R\d{4}', rng, stem)
    else:
        stem = stem + '_updated'
    return stem + suffix + ext


def prepare(epi_path, sheet_path, keep_seconds=False, time_source='created_at', audit=False, overrides=None):
    """Reads both inputs and works everything out, writing nothing: the new datasheet rows, the checks and the
    deployments. Returns a dict the writers below and fieldlog_service take.

    Raises:
        FieldlogError: an input can't be read, or no deployment can be built from them."""
    report = Report()
    if str(sheet_path).lower().endswith('.xls'):
        raise FieldlogError('Old .xls format: open it in Excel and save it as .xlsx first.')
    try:
        wb = openpyxl.load_workbook(sheet_path)
    except Exception as e:
        raise FieldlogError('Could not open the datasheet %s: %s' % (os.path.basename(str(sheet_path)), e))
    dialect = standardize_sheet(wb, report)
    ws, colidx = find_sheets(wb)
    visits = read_sheet_visits(ws, colidx, report)
    if not visits:
        raise FieldlogError('The datasheet has no visits.')
    report.info('Datasheet: %s (%d visits, sessions %s to %s)' % (
        os.path.basename(str(sheet_path)), len(visits),
        session_name(min(v['snum'] for v in visits)), session_name(max(v['snum'] for v in visits))))
    try:
        epi_rows, cols = read_epicollect(epi_path, report, overrides)
    except (OSError, UnicodeDecodeError, csv.Error) as e:
        raise FieldlogError('Could not read the Epicollect file %s: %s' % (os.path.basename(str(epi_path)), e))
    prefix = site_prefix(visits)
    new = select_new(epicollect_to_rows(epi_rows, cols, prefix, time_source, report), visits, report)
    new_visits = new_to_visits(new)
    all_visits = visits + new_visits
    last_s = max(v['snum'] for v in all_visits)
    first_s = min(v['snum'] for v in all_visits)
    check_sessions = sorted(set(v['snum'] for v in new_visits)) or [last_s]
    if audit:
        check_sessions = sorted(set(v['snum'] for v in all_visits))
    check_history(all_visits, check_sessions, report)
    intervals = build_intervals(all_visits, keep_seconds, report, check_sessions)
    if not intervals:
        raise FieldlogError('No deployment interval could be built.')
    return dict(report=report, wb=wb, ws=ws, colidx=colidx, dialect=dialect, new=new, visits=all_visits,
                new_visits=new_visits, intervals=intervals, first_s=first_s, last_s=last_s,
                target=max(iv['dep'] for iv in intervals), check_sessions=check_sessions, keep_seconds=keep_seconds)


def write_datasheets(prep, sheet_path, out_dir, output_format='original'):
    """Writes the updated datasheet beside the original — a new file, named after the new range of sessions; the
    original is never touched. Nothing when there are no new rows. Returns the paths written."""
    report, new = prep['report'], prep['new']
    if not new:
        report.info('Datasheet not rewritten (no new rows).')
        return []
    os.makedirs(out_dir, exist_ok=True)
    wb, ws, colidx, dialect = prep['wb'], prep['ws'], prep['colidx'], prep['dialect']
    outputs = []        # (workbook, sheet, columns, rows, file name suffix, description)
    if dialect is None:
        outputs.append((wb, ws, colidx, new, '', 'standard format'))
    else:
        if output_format in ('original', 'both'):
            wb_o = openpyxl.load_workbook(sheet_path)
            ws_o, colidx_o = find_original_sheet(wb_o)
            outputs.append((wb_o, ws_o, colidx_o, rows_to_original(new, ws_o, colidx_o), '', 'original format'))
        if output_format in ('standard', 'both'):
            outputs.append((wb, ws, colidx, new, '_standard', 'standard format'))
    files = []
    for wbk, wsk, colk, rowsk, suffix, what in outputs:
        append_rows(wsk, colk, rowsk)
        wbk.calculation.fullCalcOnLoad = True
        p = unique_path(os.path.join(out_dir, new_sheet_name(sheet_path, prep['first_s'], prep['last_s'], suffix)))
        wbk.save(p)
        files.append(p)
        report.info('Updated datasheet written in %s (%d rows added): %s' % (what, len(new), os.path.basename(p)))
    return files
