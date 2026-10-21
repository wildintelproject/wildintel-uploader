"""Field records — the Epicollect download of a field trip and the fieldwork datasheet turned into the updated
datasheet, the deployments of the collection and its FileTimestampLog and FieldworkLog."""
import csv
import datetime as dt
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from fastapi.testclient import TestClient

from wildintel_uploader.core import config
from wildintel_uploader.core.services import fieldlog_service as fs
from wildintel_uploader.core.services import local_folder_service as lfs
from wildintel_uploader.web.main import app

HEADERS = ["Trappers", "Session", "Date", "Time", "DateTime", "Site", "Camera", "Action", "Status", "CardBatteriesOK",
           "NumPhotos", "SD", "DateTimeOK", "Notes", "Longitude", "Latitude"]
EPI_HEADERS = ["ec5_uuid", "created_at", "1_Responsable", "2_Collection", "3_Date", "4_Location", "5_Camera", "6_Action",
               "7_Status_of_the_came", "8_Card_Batteries", "9_Number_of_pictures", "10_SD", "11_Date_and_hour",
               "12_Observation", "lat_13_Coordenates", "long_13_Coordenates"]

# The datasheet: two sites installed in R0001 and revised in R0002.
SHEET_ROWS = [
    ("Ana", "R0001", dt.date(2026, 1, 10), dt.time(9, 0), "TATR_0001", 11, "Installation", "OK", None, "1a", "pole"),
    ("Ana", "R0001", dt.date(2026, 1, 10), dt.time(10, 0), "TATR_0002", 12, "Installation", "OK", None, "2a", None),
    ("Ben", "R0002", dt.date(2026, 2, 20), dt.time(9, 30), "TATR_0001", 11, "Revision", "OK", 120, "1b", None),
    ("Ben", "R0002", dt.date(2026, 2, 20), dt.time(11, 0), "TATR_0002", 12, "Revision", "OK", 80, "2b", "tilted"),
]
# The field trip of R0003: TATR_0001 revised (its card full: no count), TATR_0003 installed — a new site. TATR_0002 is
# not visited yet.
EPI_ROWS = [
    ("u1", "2026-03-30T08:15:42.000Z", "Carla", "3", "30/03/2026", "1", "11", "Revision", "OK", "OK", "", "1c", "OK", "card full", "49.2", "20.1"),
    ("u2", "2026-03-30T10:05:00.000Z", "Carla", "3", "30/03/2026", "3", "15", "Installation", "OK", "OK", "", "3a", "OK", "", "49.3", "20.2"),
]


def _datasheet(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for trappers, session, day, time, site, camera, action, status, photos, sd, notes in SHEET_ROWS:
        ws.append([trappers, session, day, time, None, site, camera, action, status, "OK", photos, sd, "OK", notes, 20.0, 49.0])
    wb.save(path)
    return path


def _epicollect(path: Path, rows=EPI_ROWS) -> Path:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(EPI_HEADERS)
        writer.writerows(rows)
    return path


@pytest.fixture
def trip(tmp_path):
    """The two inputs, and a collections folder with the research project TATRA and its locations TATR_0001-0002."""
    root = tmp_path / "collections"
    lfs.write_research_project(root, {"acronym": "TATRA", "name": "Tatra"})
    for site in ("TATR_0001", "TATR_0002"):
        lfs.write_location(root, "TATRA", {"location_id": site, "timezone": "UTC"})
    sheet = _datasheet(tmp_path / "TatraFieldSheet_R0001-R0002.xlsx")
    epi = _epicollect(tmp_path / "form-backup.csv")
    with patch("wildintel_uploader.core.config.collections_dir", return_value=root):
        yield root, epi, sheet


def test_the_trip_closes_the_deployments_of_the_previous_collection(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert result["new_rows"] == 2
    assert result["collections"] == ["R0002"]
    by_id = {d["deployment_id"]: d for d in result["deployments"]}
    closed = by_id["R0002-TATR_0001"]
    assert closed["status"] == "closed"
    assert (closed["start"], closed["end"]) == ("2026-02-20T09:30:00Z", "2026-03-30T08:15:00Z")
    # The camera, the card and who set it up are the start visit's; a full card closes it all the same.
    assert (closed["setup_by"], closed["camera_id"], closed["card"], closed["photos_counted"]) == ("Setup: Ben | Retrieval: Carla", "11", "1b", None)
    assert closed["comments"] == "Retrieval: card full"
    # TATR_0002 is not visited yet: its camera is still out there.
    assert by_id["R0002-TATR_0002"]["status"] == "open"
    assert by_id["R0002-TATR_0002"]["comments"] == "Setup: tilted"
    assert by_id["R0003-TATR_0003"]["status"] == "open"


def test_a_site_that_is_not_a_location_yet_is_reported(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert [s["site"] for s in result["new_sites"]] == ["TATR_0003"]
    assert result["new_sites"][0]["latitude"] == 49.3
    assert any("TATR_0003" in w and "Trapper" in w for w in result["warnings"])


def test_without_locations_kept_new_sites_are_not_judged(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), None)

    assert result["locations_known"] is False and result["new_sites"] == []


def test_a_camera_at_two_sites_in_a_session_is_reported(trip, tmp_path):
    _, _, sheet = trip
    rows = [EPI_ROWS[0], EPI_ROWS[1][:6] + ("11",) + EPI_ROWS[1][7:]]
    epi = _epicollect(tmp_path / "shared.csv", rows)

    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert any("camera 11 is at 2 sites" in w for w in result["warnings"])


def test_applying_writes_the_datasheet_and_the_collection_files(trip):
    root, epi, sheet = trip
    result = fs.apply(str(epi), str(sheet), "TATRA")

    updated = sheet.parent / "TatraFieldSheet_R0001-R0003.xlsx"
    assert str(updated) in result["files"] and sheet.is_file()
    assert openpyxl.load_workbook(updated).active.max_row == 7
    collection = root / "TATRA" / "R0002"
    log = lfs.read_timestamp_log(lfs.timestamp_log_path(collection))
    assert log == [{"Deployment": "r0002-tatr_0001", "StartDate": "2026:02:20", "StartTime": "09:30:00",
                    "EndDate": "2026:03:30", "EndTime": "08:15:00"}]
    assert not (sheet.parent / "R0002_FileTimestampLog.csv").exists()  # the old eleven-column log is not written
    rows = fs.read_fieldwork_log(collection)
    assert [(r["deploymentID"], r["status"]) for r in rows] == [("R0002-TATR_0001", "closed"), ("R0002-TATR_0002", "open")]
    assert rows[0]["setupBy"] == "Setup: Ben | Retrieval: Carla" and rows[0]["cameraID"] == "11" and rows[0]["deploymentEnd"] == "2026-03-30T08:15:00Z"


def test_applying_again_keeps_the_other_rows_of_the_timestamp_log(trip):
    root, epi, sheet = trip
    collection = root / "TATRA" / "R0002"
    lfs.upsert_timestamp_log(collection, "R0002-TATR_0099", dt.datetime(2026, 2, 1), dt.datetime(2026, 3, 1))

    fs.apply(str(epi), str(sheet), "TATRA")

    ids = [r["Deployment"] for r in lfs.read_timestamp_log(lfs.timestamp_log_path(collection))]
    assert ids == ["r0002-tatr_0099", "r0002-tatr_0001"]  # in lower case, as Trapper keeps its ids


def test_a_research_project_is_needed_to_write(trip):
    _, epi, sheet = trip
    with pytest.raises(fs.FieldlogError):
        fs.apply(str(epi), str(sheet), "")
    with pytest.raises(fs.FieldlogError):
        fs.apply(str(epi), str(sheet), "NOPE")


def test_a_missing_input_is_an_error(trip, tmp_path):
    _, epi, _ = trip
    with pytest.raises(fs.FieldlogError, match="field history"):
        fs.analyze(str(epi), str(tmp_path / "missing.xlsx"))


def test_columns_named_otherwise_are_mapped_in_the_settings(trip, tmp_path):
    _, _, sheet = trip
    renamed = _epicollect(tmp_path / "renamed.csv")
    text = renamed.read_text(encoding="utf-8").replace("12_Observation", "12_Remarks", 1)
    renamed.write_text(text, encoding="utf-8")
    settings = config.Settings()
    with pytest.raises(fs.FieldlogError, match="Observation"):
        fs.analyze(str(renamed), str(sheet), "TATRA", settings)

    settings.FIELDLOG.columns = {"Notes": "12_Remarks"}
    result = fs.analyze(str(renamed), str(sheet), "TATRA", settings)

    assert result["new_rows"] == 2


def test_the_api_analyzes_and_applies(trip):
    _, epi, sheet = trip
    client = TestClient(app)
    payload = {"epicollect": str(epi), "datasheet": str(sheet), "research_project_id": "TATRA"}

    analyzed = client.post("/api/fieldlog/analyze", json=payload)
    applied = client.post("/api/fieldlog/apply", json=payload)
    last = client.get("/api/fieldlog/last").json()

    assert analyzed.status_code == 200 and analyzed.json()["collections"] == ["R0002"]
    assert applied.status_code == 200
    assert last["research_project_id"] == "TATRA" and last["datasheet"].endswith("R0001-R0003.xlsx")
    assert client.post("/api/fieldlog/analyze", json={**payload, "datasheet": "/nope.xlsx"}).status_code == 400
    assert len(client.get("/api/fieldlog/columns").json()["results"]) == 16 + len(fs.engine.EXTRA_FIELDS)


def test_the_camera_timezone_must_be_a_timezone():
    with pytest.raises(ValueError):
        config.FieldlogSettings(camera_timezone="Mars/Olympus")
    assert config.FieldlogSettings().camera_timezone == "UTC"


def test_a_research_project_and_its_locations_are_taken_from_trapper(tmp_path):
    from wildintel_uploader.core.services import sync_service
    root = tmp_path / "collections"
    lfs.write_research_project(root, {"acronym": "DONA_24", "name": "Doñana"})
    lfs.write_location(root, "DONA_24", {"location_id": "DONA_0001", "timezone": "UTC"})
    trapper = [
        {"pk": 1, "location_id": "dona_0001", "name": "a", "timezone": "Europe/Madrid", "ignore_dst": True, "latitude": 37.0, "longitude": -6.4},
        {"pk": 2, "location_id": "dona_0002", "name": "b", "timezone": "UTC", "ignore_dst": False, "latitude": 37.1, "longitude": -6.5},
    ]
    with patch("wildintel_uploader.core.services.trapper_service.list_locations", return_value=trapper):
        done = sync_service.take_locations(("u", "n", "p"), root, {"pk": 9, "name": "Doñana", "acronym": "DONA_24"})

    assert (done["created"], done["kept"]) == (["DONA_0002"], ["DONA_0001"])
    assert [l["location_id"] for l in lfs.list_locations(root, "DONA_24")] == ["DONA_0001", "DONA_0002"]
    assert lfs.list_locations(root, "DONA_24")[0]["timezone"] == "UTC"  # what was kept stays as it was


def test_a_new_site_comes_with_what_tells_it_from_a_misnamed_one(trip):
    root, epi, sheet = trip
    lfs.update_location(root, "TATRA", "TATR_0002", {"latitude": 49.3005, "longitude": 20.2})

    [site] = fs.analyze(str(epi), str(sheet), "TATRA")["new_sites"]

    assert site["site"] == "TATR_0003" and site["action"] == "Installation"
    assert (site["latitude"], site["longitude"]) == (49.3, 20.2)  # its installation's
    assert site["nearest"] == {"location_id": "TATR_0002", "distance_m": 56}
    assert site["similar"] == ["TATR_0001", "TATR_0002"]  # one character apart: misnamed?


def test_ids_one_character_apart():
    assert fs._one_apart("TATR_0004", "tatr_0005") and fs._one_apart("DONA_0022", "DONA_022")
    assert fs._one_apart("DONA_0012", "DONA_0021")  # two digits swapped
    assert not fs._one_apart("DONA_0022", "DONA_0022_B") and not fs._one_apart("TATR_0004", "TATR_0004")


def test_new_sites_are_saved_as_locations_from_the_field(trip):
    root, epi, sheet = trip
    [saved] = fs.add_locations("TATRA", [{"location_id": "tatr_0003", "latitude": 49.3, "longitude": 20.2}])

    assert saved["location_id"] == "TATR_0003" and saved["source"] == "field" and saved["timezone"] == "UTC"
    assert fs.analyze(str(epi), str(sheet), "TATRA")["new_sites"] == []
    with pytest.raises(fs.FieldlogError, match="already exists"):
        fs.add_locations("TATRA", [{"location_id": "TATR_0003", "latitude": 49.3, "longitude": 20.2}])
    with pytest.raises(fs.FieldlogError, match="valid coordinates"):
        fs.add_locations("TATRA", [{"location_id": "TATR_0009", "latitude": 120, "longitude": 20.2}])


def test_the_api_saves_new_locations(trip):
    client = TestClient(app)
    response = client.post("/api/fieldlog/add-locations", json={"research_project_id": "TATRA", "locations": [{"location_id": "TATR_0003", "latitude": 49.3, "longitude": 20.2}]})

    assert response.status_code == 200 and response.json()["results"][0]["location_id"] == "TATR_0003"


def test_a_camera_not_reached_for_a_while_is_still_a_deployment_of_the_collection_being_completed():
    import datetime as dt
    visit = lambda site, snum, day, action="Revision": {"site": site, "session": f"R{snum:04d}", "snum": snum, "when": dt.datetime(2026, 1, day), "action": action, "status": "OK", "cam": 1, "nphotos": 5}
    visits = [visit("A", 40, 1, "Installation"), visit("B", 41, 2), visit("C", 42, 3), visit("D", 43, 4)]

    found = {d["site"]: d["collection"] for d in fs._open_deployments(visits, False, 43)}

    # A, B and C will be collected on this trip (R0043): they are R0042's, like the ones collected already; D was left on it.
    assert found == {"A": "R0042", "B": "R0042", "C": "R0042", "D": "R0043"}


def test_an_installation_while_the_camera_before_was_not_collected_is_reported(trip, tmp_path):
    _, _, sheet = trip
    rows = [EPI_ROWS[0], ("u3", "2026-03-30T11:00:00.000Z", "Carla", "3", "30/03/2026", "2", "16", "Installation", "OK", "OK", "", "2c", "OK", "", "49.2", "20.1")]
    epi = _epicollect(tmp_path / "swap.csv", rows)

    warnings = fs.analyze(str(epi), str(sheet), "TATRA")["warnings"]

    assert any("TATR_0002" in w and "not collected" in w for w in warnings)


def _with_extra_columns(sheet: Path, epi: Path) -> None:
    """The setup visits of R0002 with Camtrap DP fields in the field history, and the R0003 visits with them in the field
    survey: TATR_0001's height and habitat noted otherwise when its card was collected."""
    wb = openpyxl.load_workbook(sheet)
    ws = wb.active
    extra = ["cameraHeight", "habitat", "baitUse", "deploymentTags"]
    for i, name in enumerate(extra, start=len(HEADERS) + 1):
        ws.cell(1, i).value = name
    ws.cell(4, len(HEADERS) + 1).value = 1.2  # R0002 TATR_0001
    ws.cell(4, len(HEADERS) + 2).value = "forest"
    ws.cell(4, len(HEADERS) + 3).value = "yes"
    ws.cell(4, len(HEADERS) + 4).value = "edge"
    wb.save(sheet)
    rows = [row + extra_values for row, extra_values in zip(EPI_ROWS, [("1,5", "meadow", "si", "edge|wet"), ("abc", "", "", "")])]
    with epi.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(EPI_HEADERS + extra)
        writer.writerows(rows)


def test_camtrap_fields_of_both_visits_are_kept_setup_first(trip):
    _, epi, sheet = trip
    _with_extra_columns(sheet, epi)

    result = fs.analyze(str(epi), str(sheet), "TATRA")

    closed = next(d for d in result["deployments"] if d["deployment_id"] == "R0002-TATR_0001")
    assert closed["fields"] == {"cameraHeight": 1.2, "habitat": "Setup: forest | Retrieval: meadow", "baitUse": True,
                                "deploymentTags": "edge | wet"}
    assert closed["comments"] == "Retrieval: card full; cameraHeight 1.5 (setup 1.2)"
    # A value that can't be a height is left out, and said.
    assert any("R0003 TATR_0003: cameraHeight 'abc' is not a number" in w for w in result["warnings"])


def test_camtrap_fields_reach_the_fieldwork_log_and_stay_in_the_field_history(trip):
    root, epi, sheet = trip
    _with_extra_columns(sheet, epi)

    result = fs.apply(str(epi), str(sheet), "TATRA")

    rows = fs.read_fieldwork_log(root / "TATRA" / "R0002")
    first = next(r for r in rows if r["deploymentID"] == "R0002-TATR_0001")
    assert (first["cameraHeight"], first["baitUse"], first["deploymentTags"]) == ("1.2", "true", "edge | wet")
    updated = next(f for f in result["files"] if f.endswith("R0001-R0003.xlsx"))
    ws = openpyxl.load_workbook(updated).active
    heads = [c.value for c in ws[1]]
    assert ws.cell(6, heads.index("habitat") + 1).value == "meadow"  # the R0003 visit of TATR_0001, kept for later trips


def test_a_camtrap_column_the_field_history_lacks_is_added(trip, tmp_path):
    _, _, sheet = trip
    epi = tmp_path / "with-height.csv"
    with epi.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(EPI_HEADERS + ["cameraHeight"])
        writer.writerows(row + ("0.8",) for row in EPI_ROWS)

    result = fs.apply(str(epi), str(sheet), "TATRA")

    updated = next(f for f in result["files"] if f.endswith("R0001-R0003.xlsx"))
    ws = openpyxl.load_workbook(updated).active
    heads = [c.value for c in ws[1]]
    assert "cameraHeight" in heads and ws.cell(7, heads.index("cameraHeight") + 1).value == "0.8"


def test_the_sites_not_visited_are_a_note_not_a_warning(trip):
    _, epi, sheet = trip

    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert not any("not in" in w and "visited" in w for w in result["warnings"])
    assert any("their cameras still out: TATR_0002" in n for n in result["notes"])
